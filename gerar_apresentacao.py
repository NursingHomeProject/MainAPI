#!/usr/bin/env python3
"""Gerador de apresentação PowerPoint - Casa Assistencial"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as MSO
from pptx.oxml.ns import qn
from lxml import etree

# ── Paleta de cores ────────────────────────────────────────────────────────────
NAVY     = RGBColor(0x1A, 0x2E, 0x4A)   # azul-marinho escuro
TEAL     = RGBColor(0x16, 0x79, 0x6A)   # verde-teal (cor primária)
TEAL_LT  = RGBColor(0x1A, 0xB5, 0x95)   # teal claro
TEAL_BG  = RGBColor(0xD4, 0xED, 0xE9)   # fundo teal muito claro
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY    = RGBColor(0xF4, 0xF6, 0xF8)   # fundo claro
MGRAY    = RGBColor(0x8E, 0x99, 0xAA)   # texto secundário
DGRAY    = RGBColor(0x2D, 0x3A, 0x4A)   # texto escuro
BLUE     = RGBColor(0x21, 0x4F, 0x7E)   # azul corporativo
AMBER    = RGBColor(0xF0, 0xA3, 0x0A)   # âmbar destaque
GREEN    = RGBColor(0x27, 0xAE, 0x60)   # verde sucesso
BORDER   = RGBColor(0xDD, 0xE3, 0xEB)   # borda suave

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]

# ── Helpers ────────────────────────────────────────────────────────────────────

def add_bg(slide, color):
    s = slide.shapes.add_shape(MSO.RECTANGLE, 0, 0, W, H)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s

def rect(slide, l, t, w, h, fill=None, line=None, lw=Pt(1)):
    s = slide.shapes.add_shape(MSO.RECTANGLE, l, t, w, h)
    if fill: s.fill.solid(); s.fill.fore_color.rgb = fill
    else:    s.fill.background()
    if line: s.line.color.rgb = line; s.line.width = lw
    else:    s.line.fill.background()
    return s

def rrect(slide, l, t, w, h, fill=None, line=None, lw=Pt(1)):
    s = slide.shapes.add_shape(MSO.ROUNDED_RECTANGLE, l, t, w, h)
    s.adjustments[0] = 0.06
    if fill: s.fill.solid(); s.fill.fore_color.rgb = fill
    else:    s.fill.background()
    if line: s.line.color.rgb = line; s.line.width = lw
    else:    s.line.fill.background()
    return s

def oval(slide, l, t, w, h, fill=None):
    s = slide.shapes.add_shape(MSO.OVAL, l, t, w, h)
    if fill: s.fill.solid(); s.fill.fore_color.rgb = fill
    else:    s.fill.background()
    s.line.fill.background()
    return s

def arrow_r(slide, l, t, w, h, fill=TEAL):
    s = slide.shapes.add_shape(MSO.RIGHT_ARROW, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s

def txt(slide, text, l, t, w, h, size=16, bold=False, color=DGRAY,
        align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color
    r.font.name = "Calibri"
    return tb

def shape_txt(shape, text, size=14, bold=False, color=WHITE,
              align=PP_ALIGN.CENTER):
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = "Calibri"

def title_bar(slide, title, subtitle=None):
    rect(slide, 0, 0, W, Inches(1.35), fill=NAVY)
    rect(slide, 0, 0, Inches(0.13), Inches(1.35), fill=TEAL)
    txt(slide, title, Inches(0.32), Inches(0.12), Inches(12.5), Inches(0.75),
        size=28, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    if subtitle:
        txt(slide, subtitle, Inches(0.32), Inches(0.85), Inches(12), Inches(0.4),
            size=13, color=TEAL_LT, align=PP_ALIGN.LEFT)

def section_label(slide, text, color=TEAL):
    rect(slide, Inches(0.5), Inches(7.1), Inches(0.5), Inches(0.03), fill=color)
    txt(slide, text, Inches(0.5), Inches(7.1), Inches(5), Inches(0.3),
        size=10, color=MGRAY, italic=True)

def card(slide, l, t, w, h, title, lines, hdr_color=TEAL,
         title_color=NAVY, text_color=DGRAY, bg=WHITE, icon=""):
    rrect(slide, l, t, w, h, fill=bg, line=BORDER)
    rect(slide, l, t, w, Inches(0.055), fill=hdr_color)
    header_txt = f"{icon}  {title}" if icon else title
    txt(slide, header_txt, l+Inches(0.18), t+Inches(0.1),
        w-Inches(0.36), Inches(0.45), size=13, bold=True, color=title_color)
    yo = t + Inches(0.6)
    for line in lines:
        txt(slide, line, l+Inches(0.22), yo, w-Inches(0.44), Inches(0.38),
            size=11, color=text_color)
        yo += Inches(0.37)

def add_styled_table(slide, l, t, w, rows_data, col_widths):
    n_rows = len(rows_data)
    n_cols = len(col_widths)
    total_h = Inches(0.42) * n_rows
    table = slide.shapes.add_table(n_rows, n_cols, l, t, w, total_h).table
    for ci, cw in enumerate(col_widths):
        table.columns[ci].width = cw
    for ri, row in enumerate(rows_data):
        for ci, cell_text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = cell_text
            tf = cell.text_frame
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            for para in tf.paragraphs:
                for run in para.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(11)
                    if ri == 0:
                        run.font.bold = True
                        run.font.color.rgb = WHITE
                    elif ri % 2 == 0:
                        run.font.color.rgb = DGRAY
                    else:
                        run.font.color.rgb = DGRAY
            # Row background
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            solidFill = etree.SubElement(tcPr, qn('a:solidFill'))
            srgbClr   = etree.SubElement(solidFill, qn('a:srgbClr'))
            if ri == 0:
                srgbClr.set('val', '16796A')
            elif ri % 2 == 1:
                srgbClr.set('val', 'F4F6F8')
            else:
                srgbClr.set('val', 'FFFFFF')
    return table

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — CAPA
# ══════════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(BLANK)
add_bg(s1, NAVY)

# Barra esquerda teal
rect(s1, 0, 0, Inches(0.75), H, fill=TEAL)

# Círculos decorativos (canto superior direito)
oval(s1, Inches(10.8), Inches(-1.2), Inches(4.5), Inches(4.5), fill=BLUE)
oval(s1, Inches(11.5), Inches(-0.5), Inches(3), Inches(3), fill=TEAL)
oval(s1, Inches(9.5),  Inches(5.2),  Inches(2.5), Inches(2.5), fill=RGBColor(0x1A,0x2E,0x6A))

# Cruz médica estilizada
rect(s1, Inches(1.3), Inches(1.4), Inches(1.4), Inches(0.45), fill=WHITE)
rect(s1, Inches(1.65), Inches(1.1), Inches(0.7), Inches(1.1), fill=WHITE)

# Tag
txt(s1, "PLATAFORMA DE GESTÃO ASSISTENCIAL",
    Inches(1.3), Inches(0.7), Inches(10), Inches(0.45),
    size=12, bold=False, color=TEAL_LT)

# Título principal
txt(s1, "Casa Assistencial",
    Inches(1.3), Inches(2.5), Inches(10.5), Inches(1.5),
    size=54, bold=True, color=WHITE)

# Linha decorativa
rect(s1, Inches(1.3), Inches(4.15), Inches(7), Inches(0.045), fill=TEAL_LT)

# Subtítulo
txt(s1, "Plataforma de Gestão Assistencial em Nuvem",
    Inches(1.3), Inches(4.3), Inches(10), Inches(0.65),
    size=22, color=TEAL_LT)

# Descritivo
txt(s1, "Arquitetura · Infraestrutura · Custos de Produção",
    Inches(1.3), Inches(5.1), Inches(10), Inches(0.45),
    size=14, color=MGRAY)

# Rodapé
txt(s1, "Versão 1.0  ·  2026",
    Inches(1.3), Inches(6.8), Inches(5), Inches(0.4),
    size=11, color=MGRAY)

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — OBJETIVO DO SISTEMA
# ══════════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(BLANK)
add_bg(s2, LGRAY)
title_bar(s2, "Objetivo do Sistema",
          "Por que este sistema foi criado e qual problema ele resolve")

# 4 cards 2x2
card_data = [
    ("Controle de Pacientes", [
        "Cadastro completo de cada residente",
        "Histórico de cuidados e observações",
        "Situação de alta ou internação"
    ], TEAL, "👤"),
    ("Medicamentos e Insumos", [
        "Controle de fraldas, remédios e materiais",
        "Registro de cada administração",
        "Histórico de consumo por paciente"
    ], BLUE, "💊"),
    ("Alertas Automáticos", [
        "Notificação de estoque abaixo do mínimo",
        "Alerta de dose atrasada ou esquecida",
        "Identificação de padrões irregulares"
    ], AMBER, "🔔"),
    ("Gestão Operacional", [
        "Dashboard com indicadores em tempo real",
        "Redução de erros de administração",
        "Decisões baseadas em dados confiáveis"
    ], GREEN, "📊"),
]
positions = [
    (Inches(0.5),  Inches(1.55)),
    (Inches(6.8),  Inches(1.55)),
    (Inches(0.5),  Inches(4.4)),
    (Inches(6.8),  Inches(4.4)),
]
cw, ch = Inches(6.1), Inches(2.65)
for (l,t), (title, lines, hdr, icon) in zip(positions, card_data):
    card(s2, l, t, cw, ch, title, lines, hdr_color=hdr, icon=icon)

section_label(s2, "Casa Assistencial · Objetivo do Sistema")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — COMO O USUÁRIO UTILIZA O SISTEMA
# ══════════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(BLANK)
add_bg(s3, LGRAY)
title_bar(s3, "Como o Usuário Acessa o Sistema",
          "O sistema funciona pelo navegador — sem instalar nada")

devices = [
    ("💻", "Computador", [
        "Acesso pelo navegador",
        "(Chrome, Firefox, Edge)",
        "",
        "Ideal para trabalho",
        "administrativo completo",
    ], TEAL),
    ("📱", "Tablet", [
        "Acesso pelo navegador",
        "do tablet (Safari, Chrome)",
        "",
        "Perfeito para consulta",
        "em campo e rondas",
    ], BLUE),
    ("📲", "Celular", [
        "Funciona como aplicativo",
        "(Tecnologia PWA)",
        "",
        "Instala na tela inicial",
        "sem precisar de loja",
    ], AMBER),
]

for i, (icon, device, lines, color) in enumerate(devices):
    l = Inches(0.6) + i * Inches(4.2)
    t = Inches(1.6)
    cw2 = Inches(3.9)
    # Card principal
    rrect(s3, l, t, cw2, Inches(5.5), fill=WHITE, line=BORDER)
    rect(s3,  l, t, cw2, Inches(0.07), fill=color)
    # Ícone grande
    txt(s3, icon, l, t+Inches(0.25), cw2, Inches(1.0),
        size=42, align=PP_ALIGN.CENTER, color=DGRAY)
    # Nome do dispositivo
    txt(s3, device, l, t+Inches(1.25), cw2, Inches(0.5),
        size=18, bold=True, color=color, align=PP_ALIGN.CENTER)
    # Linhas de detalhe
    yo = t + Inches(1.85)
    for line in lines:
        txt(s3, line, l+Inches(0.2), yo, cw2-Inches(0.4), Inches(0.35),
            size=12, color=DGRAY, align=PP_ALIGN.CENTER)
        yo += Inches(0.33)

# Nota PWA
note_box = rrect(s3, Inches(3.0), Inches(6.75), Inches(7.3), Inches(0.55),
                 fill=TEAL_BG, line=TEAL)
txt(s3, "✓  PWA (Progressive Web App): o sistema pode ser instalado como app no celular sem precisar de Google Play ou App Store",
    Inches(3.2), Inches(6.8), Inches(7.0), Inches(0.45),
    size=11, color=TEAL, italic=True)

section_label(s3, "Casa Assistencial · Acesso ao Sistema")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — VISÃO GERAL DA ARQUITETURA
# ══════════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(BLANK)
add_bg(s4, LGRAY)
title_bar(s4, "Visão Geral da Arquitetura",
          "Como as partes do sistema se organizam e se conectam")

arch_items = [
    ("🌐", "Internet",    "Conexão do\nusuário",               MGRAY, Inches(0.5)),
    ("🖥", "Frontend",    "Telas e\ninterface visual",          TEAL,  Inches(2.85)),
    ("⚙",  "Backend",     "Regras e\ncálculos",                 BLUE,  Inches(5.2)),
    ("🗄", "Banco de\nDados", "Armazenamento\nde informações", NAVY,  Inches(7.55)),
    ("☁",  "Nuvem OCI",  "Servidor\nremoto seguro",            AMBER, Inches(9.9)),
]

for (icon, name, desc, color, l) in arch_items:
    t = Inches(2.0)
    cw3 = Inches(2.1)
    ch3 = Inches(3.2)
    # Caixa principal
    rrect(s4, l, t, cw3, ch3, fill=color, line=None)
    # Ícone
    txt(s4, icon, l, t+Inches(0.2), cw3, Inches(0.9),
        size=34, color=WHITE, align=PP_ALIGN.CENTER)
    # Nome
    txt(s4, name, l, t+Inches(1.1), cw3, Inches(0.7),
        size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # Descrição
    txt(s4, desc, l, t+Inches(1.8), cw3, Inches(1.0),
        size=11, color=WHITE, align=PP_ALIGN.CENTER)

# Setas entre caixas
arrow_positions = [
    Inches(2.63), Inches(4.98), Inches(7.33), Inches(9.68)
]
for al in arrow_positions:
    arrow_r(s4, al, Inches(3.3), Inches(0.2), Inches(0.25), fill=TEAL)

# Legenda do domínio
dom_box = rrect(s4, Inches(2.85), Inches(5.5), Inches(7.5), Inches(0.6),
                fill=WHITE, line=TEAL)
txt(s4, "🌍  Domínio:  sistema.minhainstituicao.com.br  →  aponta para o servidor na nuvem via DNS",
    Inches(3.0), Inches(5.55), Inches(7.2), Inches(0.5),
    size=12, color=TEAL, align=PP_ALIGN.CENTER)

# Rótulo externo
rect(s4, Inches(0.5), Inches(1.6), Inches(12.4), Inches(3.9), fill=None, line=BORDER)
txt(s4, "AMBIENTE DA NUVEM (OCI)",
    Inches(2.85), Inches(1.45), Inches(7), Inches(0.35),
    size=11, bold=True, color=MGRAY)

section_label(s4, "Casa Assistencial · Arquitetura Geral")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — O QUE É O FRONTEND
# ══════════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(BLANK)
add_bg(s5, LGRAY)
title_bar(s5, "Frontend — A Parte Visual do Sistema",
          "Tudo que o usuário vê e com que interage na tela")

# Coluna esquerda — grande card visual
rrect(s5, Inches(0.5), Inches(1.55), Inches(5.5), Inches(5.6), fill=TEAL_BG, line=TEAL)
rect(s5, Inches(0.5), Inches(1.55), Inches(5.5), Inches(0.07), fill=TEAL)
txt(s5, "🖥", Inches(0.5), Inches(1.75), Inches(5.5), Inches(1.5),
    size=70, color=TEAL, align=PP_ALIGN.CENTER)
txt(s5, "Interface do Usuário", Inches(0.5), Inches(3.35), Inches(5.5), Inches(0.55),
    size=18, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
txt(s5, "Construído com React + TypeScript\nDesign moderno e responsivo",
    Inches(0.5), Inches(3.95), Inches(5.5), Inches(0.8),
    size=13, color=DGRAY, align=PP_ALIGN.CENTER, italic=True)

# Coluna direita — bullets
right_items = [
    ("Telas e Formulários",
     "Cadastro de pacientes, prescrições e movimentações de estoque"),
    ("Dashboard",
     "Painel com indicadores e resumo operacional em tempo real"),
    ("Alertas Visuais",
     "Notificações de estoque baixo e doses atrasadas com destaque"),
    ("Responsivo",
     "Adapta automaticamente ao tamanho do computador, tablet ou celular"),
]
yo = Inches(1.55)
for title_r, desc_r in right_items:
    rrect(s5, Inches(6.3), yo, Inches(6.5), Inches(1.2), fill=WHITE, line=BORDER)
    rect(s5, Inches(6.3), yo, Inches(0.07), Inches(1.2), fill=TEAL)
    txt(s5, title_r, Inches(6.5), yo+Inches(0.1), Inches(6.2), Inches(0.4),
        size=14, bold=True, color=NAVY)
    txt(s5, desc_r, Inches(6.5), yo+Inches(0.5), Inches(6.2), Inches(0.6),
        size=12, color=DGRAY)
    yo += Inches(1.35)

section_label(s5, "Casa Assistencial · Frontend")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — O QUE É O BACKEND
# ══════════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(BLANK)
add_bg(s6, LGRAY)
title_bar(s6, "Backend — O Cérebro do Sistema",
          "Processa as regras, validações e cálculos do negócio")

rrect(s6, Inches(7.3), Inches(1.55), Inches(5.5), Inches(5.6), fill=RGBColor(0xE8,0xF0,0xF8), line=BLUE)
rect(s6, Inches(7.3), Inches(1.55), Inches(5.5), Inches(0.07), fill=BLUE)
txt(s6, "⚙", Inches(7.3), Inches(1.75), Inches(5.5), Inches(1.5),
    size=70, color=BLUE, align=PP_ALIGN.CENTER)
txt(s6, "Motor de Regras e Cálculos", Inches(7.3), Inches(3.35), Inches(5.5), Inches(0.55),
    size=18, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
txt(s6, "Construído com FastAPI + Python\nRápido, seguro e escalável",
    Inches(7.3), Inches(3.95), Inches(5.5), Inches(0.8),
    size=13, color=DGRAY, align=PP_ALIGN.CENTER, italic=True)

right_items6 = [
    ("Regras de Negócio",
     "Define o que pode ou não acontecer (ex: dose já administrada não pode ser duplicada)"),
    ("Cálculo de Consumo",
     "Calcula automaticamente se o consumo real bate com a prescrição"),
    ("Autenticação e Segurança",
     "Garante que só usuários autorizados acessam as informações"),
    ("Geração de Alertas",
     "Identifica situações de risco e gera alertas automáticos"),
]
yo = Inches(1.55)
for title_r, desc_r in right_items6:
    rrect(s6, Inches(0.5), yo, Inches(6.5), Inches(1.2), fill=WHITE, line=BORDER)
    rect(s6, Inches(0.5), yo, Inches(0.07), Inches(1.2), fill=BLUE)
    txt(s6, title_r, Inches(0.7), yo+Inches(0.1), Inches(6.2), Inches(0.4),
        size=14, bold=True, color=NAVY)
    txt(s6, desc_r, Inches(0.7), yo+Inches(0.5), Inches(6.2), Inches(0.6),
        size=12, color=DGRAY)
    yo += Inches(1.35)

section_label(s6, "Casa Assistencial · Backend")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — O QUE É O BANCO DE DADOS
# ══════════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(BLANK)
add_bg(s7, LGRAY)
title_bar(s7, "Banco de Dados — A Memória do Sistema",
          "Onde todas as informações ficam guardadas com segurança")

# Analogia visual central
rrect(s7, Inches(4.5), Inches(1.6), Inches(4.3), Inches(4.5), fill=RGBColor(0xEA, 0xF0, 0xF8), line=NAVY)
rect(s7, Inches(4.5), Inches(1.6), Inches(4.3), Inches(0.07), fill=NAVY)
txt(s7, "🗄", Inches(4.5), Inches(1.8), Inches(4.3), Inches(1.3),
    size=64, color=NAVY, align=PP_ALIGN.CENTER)
txt(s7, "PostgreSQL", Inches(4.5), Inches(3.15), Inches(4.3), Inches(0.5),
    size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

# Tabelas armazenadas
db_tables = ["Pacientes", "Prescrições", "Medicamentos", "Estoque", "Alertas", "Usuários"]
txt(s7, "Informações armazenadas:", Inches(4.7), Inches(3.7), Inches(4.0), Inches(0.35),
    size=11, bold=True, color=DGRAY, align=PP_ALIGN.CENTER)
yo_db = Inches(4.1)
for i, t_name in enumerate(db_tables):
    col = i % 2
    rrect(s7, Inches(4.7) + col*Inches(2.0), yo_db, Inches(1.85), Inches(0.38),
          fill=NAVY, line=None)
    txt(s7, t_name, Inches(4.7) + col*Inches(2.0), yo_db,
        Inches(1.85), Inches(0.38), size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    if col == 1: yo_db += Inches(0.45)

# Cards laterais
left_cards = [
    ("🔒 Segurança", "Dados protegidos com\ncriptografia e backups\nautomáticos"),
    ("📋 Organização", "Cada tipo de informação\nem sua própria tabela\nestruturada"),
]
right_cards = [
    ("⚡ Velocidade", "Consultas rápidas mesmo\ncom grande volume\nde registros"),
    ("📦 Confiabilidade", "PostgreSQL: banco usado\nby grandes empresas\nno mundo todo"),
]
for i, (title_c, desc_c) in enumerate(left_cards):
    t_pos = Inches(1.6) + i * Inches(2.75)
    rrect(s7, Inches(0.4), t_pos, Inches(3.8), Inches(2.4), fill=WHITE, line=BORDER)
    rect(s7, Inches(0.4), t_pos, Inches(0.07), Inches(2.4), fill=NAVY)
    txt(s7, title_c, Inches(0.6), t_pos+Inches(0.1), Inches(3.5), Inches(0.45),
        size=13, bold=True, color=NAVY)
    txt(s7, desc_c, Inches(0.6), t_pos+Inches(0.55), Inches(3.5), Inches(1.6),
        size=12, color=DGRAY)

for i, (title_c, desc_c) in enumerate(right_cards):
    t_pos = Inches(1.6) + i * Inches(2.75)
    rrect(s7, Inches(9.1), t_pos, Inches(3.8), Inches(2.4), fill=WHITE, line=BORDER)
    rect(s7, Inches(9.1), t_pos, Inches(0.07), Inches(2.4), fill=NAVY)
    txt(s7, title_c, Inches(9.3), t_pos+Inches(0.1), Inches(3.5), Inches(0.45),
        size=13, bold=True, color=NAVY)
    txt(s7, desc_c, Inches(9.3), t_pos+Inches(0.55), Inches(3.5), Inches(1.6),
        size=12, color=DGRAY)

section_label(s7, "Casa Assistencial · Banco de Dados")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — COMO AS PARTES SE CONECTAM
# ══════════════════════════════════════════════════════════════════════════════
s8 = prs.slides.add_slide(BLANK)
add_bg(s8, LGRAY)
title_bar(s8, "Como as Partes se Conectam",
          "O caminho percorrido desde o clique do usuário até a resposta")

flow_steps = [
    ("👤", "Usuário",    "Clica ou\ndigita algo",        TEAL,  Inches(0.4)),
    ("🖥", "Frontend",   "Mostra as\ntelas e recebe\nações", BLUE, Inches(3.1)),
    ("⚙", "Backend",    "Processa,\ncalcula e\nvalida",   NAVY, Inches(5.8)),
    ("🗄", "Banco de\nDados", "Salva ou\nbusca as\ninformações", RGBColor(0x2C,0x5F,0x2E), Inches(8.5)),
    ("📲", "Resultado",  "Resposta\nexibida na\ntela",     RGBColor(0x8E,0x44,0xAD), Inches(11.2)),
]

bw = Inches(2.0)
bh = Inches(3.8)
bt = Inches(2.2)
for (icon, name, desc, color, l) in flow_steps:
    rrect(s8, l, bt, bw, bh, fill=color)
    txt(s8, icon, l, bt+Inches(0.2), bw, Inches(0.9), size=32, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s8, name, l, bt+Inches(1.1), bw, Inches(0.55), size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s8, desc, l, bt+Inches(1.7), bw, Inches(1.8), size=11, color=WHITE, align=PP_ALIGN.CENTER)

# Setas
for al in [Inches(2.42), Inches(5.12), Inches(7.82), Inches(10.52)]:
    arrow_r(s8, al, Inches(3.85), Inches(0.65), Inches(0.3), fill=AMBER)

# Nota
rrect(s8, Inches(2.0), Inches(6.3), Inches(9.3), Inches(0.7), fill=TEAL_BG, line=TEAL)
txt(s8, "💡  Tudo isso acontece em frações de segundo. O usuário só vê a tela atualizada com as informações corretas.",
    Inches(2.2), Inches(6.38), Inches(9.0), Inches(0.55), size=12, color=TEAL, italic=True)

section_label(s8, "Casa Assistencial · Fluxo de Dados")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — O QUE É COMPUTAÇÃO EM NUVEM
# ══════════════════════════════════════════════════════════════════════════════
s9 = prs.slides.add_slide(BLANK)
add_bg(s9, LGRAY)
title_bar(s9, "O que é Computação em Nuvem?",
          "Uma explicação simples para um conceito poderoso")

# Analogia central
rrect(s9, Inches(0.5), Inches(1.6), Inches(12.3), Inches(2.0), fill=TEAL_BG, line=TEAL)
txt(s9, "☁", Inches(0.7), Inches(1.65), Inches(1.5), Inches(1.5),
    size=48, color=TEAL, align=PP_ALIGN.CENTER)
txt(s9, "Analogia simples:",
    Inches(2.2), Inches(1.65), Inches(10), Inches(0.45), size=13, bold=True, color=TEAL)
txt(s9,
    "Assim como a energia elétrica vem de uma usina distante e chega em sua casa pela tomada, "
    "a computação em nuvem entrega poder de processamento e armazenamento pela internet — "
    "sem precisar ter um servidor físico na própria instituição.",
    Inches(2.2), Inches(2.1), Inches(10.3), Inches(1.2), size=13, color=DGRAY)

# 3 cards comparativos
comp = [
    ("🏢 Servidor Próprio",
     ["Comprar hardware caro", "Manter equipe técnica", "Risco de falha físca",
      "Difícil de ampliar", "Custo fixo alto"], MGRAY),
    ("☁ Nuvem",
     ["Pagar só pelo que usa", "Sem hardware para comprar", "Alta disponibilidade",
      "Amplia com um clique", "Custo sob medida"], TEAL),
    ("✅ Resultado",
     ["Mais barato para iniciar", "Menos risco operacional", "Escalabilidade simples",
      "Atualizações automáticas", "Suporte 24/7 incluído"], GREEN),
]
cw9 = Inches(3.9)
for i, (title_c, items, color) in enumerate(comp):
    l = Inches(0.5) + i * Inches(4.2)
    rrect(s9, l, Inches(3.85), cw9, Inches(3.3), fill=WHITE, line=BORDER)
    rect(s9, l, Inches(3.85), cw9, Inches(0.07), fill=color)
    txt(s9, title_c, l+Inches(0.15), Inches(3.95), cw9-Inches(0.3), Inches(0.45),
        size=13, bold=True, color=color if color != MGRAY else DGRAY)
    yo9 = Inches(4.45)
    for item in items:
        txt(s9, f"• {item}", l+Inches(0.2), yo9, cw9-Inches(0.4), Inches(0.35),
            size=11, color=DGRAY)
        yo9 += Inches(0.38)

section_label(s9, "Casa Assistencial · Computação em Nuvem")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — O QUE É A OCI
# ══════════════════════════════════════════════════════════════════════════════
s10 = prs.slides.add_slide(BLANK)
add_bg(s10, LGRAY)
title_bar(s10, "Oracle Cloud Infrastructure (OCI)",
          "A plataforma de nuvem onde o sistema está hospedado")

# Banner Oracle
rrect(s10, Inches(0.5), Inches(1.6), Inches(12.3), Inches(1.0), fill=NAVY, line=None)
txt(s10, "☁  Oracle Cloud Infrastructure  —  Uma das maiores nuvens do mundo",
    Inches(0.7), Inches(1.72), Inches(11.8), Inches(0.7),
    size=16, bold=True, color=WHITE)

oci_cards = [
    ("🆓 Always Free", [
        "Até 4 processadores (OCPU)",
        "Até 24 GB de memória RAM",
        "200 GB de armazenamento",
        "10 TB de transferência/mês",
        "Banco de dados grátis incluído",
        "Nunca expira",
    ], TEAL),
    ("💰 Custo Controlado", [
        "Paga somente pelo que usar",
        "Sem contrato de fidelidade",
        "Pode cancelar a qualquer momento",
        "Previsibilidade de gastos",
        "Créditos iniciais disponíveis",
        "Upgrade simples quando necessário",
    ], BLUE),
    ("🔒 Segurança", [
        "Data centers certificados",
        "Backups automáticos",
        "Proteção contra ataques DDoS",
        "Criptografia dos dados",
        "Conformidade com padrões globais",
        "Monitoramento 24 horas",
    ], NAVY),
]

cw10 = Inches(3.9)
for i, (title_c, items, color) in enumerate(oci_cards):
    l = Inches(0.5) + i * Inches(4.2)
    rrect(s10, l, Inches(2.85), cw10, Inches(4.2), fill=WHITE, line=BORDER)
    rect(s10, l, Inches(2.85), cw10, Inches(0.07), fill=color)
    txt(s10, title_c, l+Inches(0.15), Inches(2.95), cw10-Inches(0.3), Inches(0.45),
        size=13, bold=True, color=color)
    yo10 = Inches(3.45)
    for item in items:
        txt(s10, f"• {item}", l+Inches(0.2), yo10, cw10-Inches(0.4), Inches(0.38),
            size=11, color=DGRAY)
        yo10 += Inches(0.37)

section_label(s10, "Casa Assistencial · Oracle Cloud Infrastructure")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — O QUE É UM DOMÍNIO
# ══════════════════════════════════════════════════════════════════════════════
s11 = prs.slides.add_slide(BLANK)
add_bg(s11, LGRAY)
title_bar(s11, "O que é um Domínio?",
          "O endereço personalizado pelo qual o sistema é acessado")

# Analogia
rrect(s11, Inches(0.5), Inches(1.6), Inches(12.3), Inches(1.55), fill=TEAL_BG, line=TEAL)
txt(s11, "📮", Inches(0.7), Inches(1.65), Inches(1.4), Inches(1.1),
    size=40, color=TEAL, align=PP_ALIGN.CENTER)
txt(s11, "Analogia:",
    Inches(2.1), Inches(1.65), Inches(10), Inches(0.35), size=12, bold=True, color=TEAL)
txt(s11, "O domínio é como o endereço de uma casa. Em vez de memorizar um número complicado "
         "(ex: 192.168.1.45), você usa um nome fácil como  sistema.suainstituicao.com.br.",
    Inches(2.1), Inches(2.0), Inches(10.3), Inches(0.9), size=13, color=DGRAY)

# Exemplos de domínio
txt(s11, "Exemplos de domínio:", Inches(0.5), Inches(3.35), Inches(12), Inches(0.4),
    size=13, bold=True, color=NAVY)
domain_examples = [
    ("🌐  sistema.casaassistencial.org.br", "Domínio .org.br para organizações sem fins lucrativos",    TEAL),
    ("🌐  gestao.casaassistencial.com.br",  "Domínio .com.br mais comum para empresas brasileiras",     BLUE),
    ("🌐  casaassistencial.com",            "Domínio internacional .com (aceito no Brasil também)",      NAVY),
]
yo11 = Inches(3.85)
for domain, desc, color in domain_examples:
    rrect(s11, Inches(0.5), yo11, Inches(12.3), Inches(0.8), fill=WHITE, line=BORDER)
    rect(s11, Inches(0.5), yo11, Inches(0.07), Inches(0.8), fill=color)
    txt(s11, domain, Inches(0.75), yo11+Inches(0.05), Inches(6), Inches(0.4),
        size=14, bold=True, color=color)
    txt(s11, desc, Inches(0.75), yo11+Inches(0.42), Inches(11.5), Inches(0.32),
        size=11, color=MGRAY)
    yo11 += Inches(0.92)

# Custo
rrect(s11, Inches(0.5), Inches(6.9), Inches(12.3), Inches(0.4), fill=AMBER, line=None)
txt(s11, "💰  Custo típico: domínio .com.br ≈ R$ 40/ano via Registro.br  |  domínio .com ≈ R$ 60/ano via Cloudflare",
    Inches(0.7), Inches(6.93), Inches(12.0), Inches(0.35), size=12, bold=True, color=NAVY)

section_label(s11, "Casa Assistencial · Domínio")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 12 — DNS E HTTPS
# ══════════════════════════════════════════════════════════════════════════════
s12 = prs.slides.add_slide(BLANK)
add_bg(s12, LGRAY)
title_bar(s12, "DNS e HTTPS — Endereço e Segurança",
          "Como o sistema encontra o servidor certo e protege os dados")

# DNS
rrect(s12, Inches(0.5), Inches(1.6), Inches(5.9), Inches(5.5), fill=WHITE, line=BORDER)
rect(s12, Inches(0.5), Inches(1.6), Inches(5.9), Inches(0.07), fill=BLUE)
txt(s12, "🔍  DNS — Sistema de Nomes", Inches(0.65), Inches(1.7), Inches(5.6), Inches(0.5),
    size=16, bold=True, color=NAVY)
txt(s12, "O que é:", Inches(0.65), Inches(2.25), Inches(5.6), Inches(0.35),
    size=12, bold=True, color=BLUE)
txt(s12, "Funciona como a agenda telefônica da internet. "
         "Quando você digita um endereço, o DNS encontra o número (IP) "
         "correspondente ao servidor correto.",
    Inches(0.65), Inches(2.6), Inches(5.6), Inches(1.1), size=12, color=DGRAY)
txt(s12, "Como funciona:", Inches(0.65), Inches(3.75), Inches(5.6), Inches(0.35),
    size=12, bold=True, color=BLUE)
dns_steps = [
    "1. Você digita:  sistema.instituicao.com.br",
    "2. O DNS consulta:  onde está esse endereço?",
    "3. Responde:  está no servidor 129.148.49.83",
    "4. Seu navegador abre o sistema",
]
yo12 = Inches(4.15)
for step in dns_steps:
    txt(s12, step, Inches(0.75), yo12, Inches(5.4), Inches(0.38), size=11, color=DGRAY)
    yo12 += Inches(0.38)
rrect(s12, Inches(0.65), Inches(6.35), Inches(5.5), Inches(0.45), fill=TEAL_BG, line=TEAL)
txt(s12, "✓  Cloudflare oferece DNS gratuito e muito rápido",
    Inches(0.8), Inches(6.4), Inches(5.3), Inches(0.35), size=11, color=TEAL, bold=True)

# HTTPS
rrect(s12, Inches(6.9), Inches(1.6), Inches(5.9), Inches(5.5), fill=WHITE, line=BORDER)
rect(s12, Inches(6.9), Inches(1.6), Inches(5.9), Inches(0.07), fill=GREEN)
txt(s12, "🔒  HTTPS — Conexão Segura", Inches(7.05), Inches(1.7), Inches(5.6), Inches(0.5),
    size=16, bold=True, color=NAVY)
txt(s12, "O que é:", Inches(7.05), Inches(2.25), Inches(5.6), Inches(0.35),
    size=12, bold=True, color=GREEN)
txt(s12, "O HTTPS garante que os dados trocados entre o usuário e o servidor "
         "sejam criptografados — ninguém consegue interceptar senhas, "
         "prontuários ou informações de pacientes.",
    Inches(7.05), Inches(2.6), Inches(5.6), Inches(1.1), size=12, color=DGRAY)
txt(s12, "Vantagens:", Inches(7.05), Inches(3.75), Inches(5.6), Inches(0.35),
    size=12, bold=True, color=GREEN)
https_items = [
    "🔐  Dados criptografados de ponta a ponta",
    "✅  Cadeado verde no navegador (confiança)",
    "🏆  Melhor posicionamento em buscas",
    "🚫  Proteção contra interceptação de dados",
]
yo12b = Inches(4.15)
for item in https_items:
    txt(s12, item, Inches(7.15), yo12b, Inches(5.5), Inches(0.38), size=11, color=DGRAY)
    yo12b += Inches(0.38)
rrect(s12, Inches(7.05), Inches(6.35), Inches(5.5), Inches(0.45), fill=RGBColor(0xD5,0xF5,0xE3), line=GREEN)
txt(s12, "✓  Cloudflare oferece certificado HTTPS completamente grátis",
    Inches(7.2), Inches(6.4), Inches(5.3), Inches(0.35), size=11, color=GREEN, bold=True)

section_label(s12, "Casa Assistencial · DNS e HTTPS")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 13 — FLUXO COMPLETO DE ACESSO
# ══════════════════════════════════════════════════════════════════════════════
s13 = prs.slides.add_slide(BLANK)
add_bg(s13, LGRAY)
title_bar(s13, "Fluxo Completo de Acesso",
          "Da digitação do endereço até a exibição das informações na tela")

flow13 = [
    ("1", "Usuário digita\no endereço",    "Ex: sistema.\ninstituicao.com.br",  TEAL,  Inches(0.35)),
    ("2", "DNS encontra\no servidor",      "Aponta para\nIP na OCI",            BLUE,  Inches(2.7)),
    ("3", "HTTPS\ncriptografa",            "Conexão segura\nestabelecida",       GREEN, Inches(5.05)),
    ("4", "Servidor\nresponde",            "Frontend\ncarregado",               NAVY,  Inches(7.4)),
    ("5", "Dados\nbuscados",              "Backend + Banco\nde Dados",          RGBColor(0x8E,0x44,0xAD), Inches(9.75)),
    ("6", "Tela\nexibida",               "Informações\nna tela",               AMBER, Inches(12.1)),
]

bw13 = Inches(2.1)
bh13 = Inches(3.3)
bt13 = Inches(2.3)

for (num, title_f, desc_f, color, l) in flow13:
    if l + bw13 > W - Inches(0.1):
        continue
    rrect(s13, l, bt13, bw13, bh13, fill=color)
    # Número
    oval(s13, l + Inches(0.7), bt13 - Inches(0.25), Inches(0.7), Inches(0.7), fill=WHITE)
    txt(s13, num, l+Inches(0.7), bt13-Inches(0.3), Inches(0.7), Inches(0.7),
        size=14, bold=True, color=color, align=PP_ALIGN.CENTER)
    txt(s13, title_f, l, bt13+Inches(0.35), bw13, Inches(0.85),
        size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s13, desc_f, l, bt13+Inches(1.2), bw13, Inches(1.2),
        size=11, color=WHITE, align=PP_ALIGN.CENTER)

# Setas
for i, (_, _, _, _, l) in enumerate(flow13[:-1]):
    if l + bw13 + Inches(0.4) < Inches(12.0):
        arrow_r(s13, l + bw13 + Inches(0.05), bt13 + Inches(1.4), Inches(0.55), Inches(0.28), fill=AMBER)

# Nota de tempo
rrect(s13, Inches(2.0), Inches(6.1), Inches(9.3), Inches(0.65), fill=TEAL_BG, line=TEAL)
txt(s13, "⏱  Todo este processo acontece em menos de 1 segundo. O usuário vê as informações quase instantaneamente.",
    Inches(2.2), Inches(6.18), Inches(9.0), Inches(0.5), size=12, color=TEAL, italic=True)

section_label(s13, "Casa Assistencial · Fluxo de Acesso")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 14 — INFRAESTRUTURA RECOMENDADA
# ══════════════════════════════════════════════════════════════════════════════
s14 = prs.slides.add_slide(BLANK)
add_bg(s14, LGRAY)
title_bar(s14, "Infraestrutura Recomendada para o Projeto",
          "Componentes necessários para o sistema funcionar em produção")

infra = [
    ("🖥 Frontend",       "Interface visual", "Servidor Nginx\nnuvem OCI",          "Incluso na nuvem", TEAL),
    ("⚙ Backend",         "Motor do sistema", "FastAPI + Python\nservidor OCI",       "Incluso na nuvem", BLUE),
    ("🗄 Banco de Dados",  "Armazenamento",   "PostgreSQL\nno servidor OCI",          "Incluso na nuvem", NAVY),
    ("🌐 Domínio",         "Endereço web",    "Registro.br (.com.br)\nou Cloudflare",  "≈ R$ 40–60/ano",   AMBER),
    ("🔒 DNS + HTTPS",     "Segurança",       "Cloudflare Free\ncertificado grátis",   "GRATUITO",         GREEN),
    ("📧 E-mail (futuro)", "Notificações",    "SendGrid ou\nSES Amazon",              "≈ R$ 20–50/mês",   MGRAY),
]

cw14 = Inches(3.9)
for i, (name, role, tech, cost, color) in enumerate(infra):
    row = i // 3
    col = i % 3
    l14 = Inches(0.4) + col * Inches(4.28)
    t14 = Inches(1.6) + row * Inches(2.7)
    rrect(s14, l14, t14, cw14, Inches(2.45), fill=WHITE, line=BORDER)
    rect(s14, l14, t14, cw14, Inches(0.07), fill=color)
    txt(s14, name, l14+Inches(0.15), t14+Inches(0.1), cw14-Inches(0.3), Inches(0.45),
        size=14, bold=True, color=color if color != MGRAY else DGRAY)
    txt(s14, role, l14+Inches(0.15), t14+Inches(0.55), cw14-Inches(0.3), Inches(0.35),
        size=11, color=MGRAY, italic=True)
    txt(s14, tech, l14+Inches(0.15), t14+Inches(0.9), cw14-Inches(0.3), Inches(0.75),
        size=12, color=DGRAY)
    cost_color = GREEN if "GRATUITO" in cost or "Incluso" in cost else AMBER
    rrect(s14, l14+Inches(0.1), t14+Inches(1.75), cw14-Inches(0.2), Inches(0.4),
          fill=TEAL_BG if "GRATUITO" in cost or "Incluso" in cost else RGBColor(0xFE, 0xF9, 0xE7), line=None)
    txt(s14, f"💰 {cost}", l14+Inches(0.15), t14+Inches(1.78), cw14-Inches(0.3), Inches(0.35),
        size=11, bold=True, color=cost_color, align=PP_ALIGN.CENTER)

section_label(s14, "Casa Assistencial · Infraestrutura")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 15 — PESQUISA DE CUSTOS REAIS
# ══════════════════════════════════════════════════════════════════════════════
s15 = prs.slides.add_slide(BLANK)
add_bg(s15, LGRAY)
title_bar(s15, "Pesquisa de Custos Reais",
          "Valores levantados em maio de 2026 · Cotação: US$ 1 = R$ 5,70")

cost_cards = [
    ("🌐 Domínio .com.br", "Registro.br",
     "R$ 40,00/ano", "≈ R$ 3,33/mês",
     "Registro oficial de domínios brasileiros.\nPreço fixo sem taxas adicionais.",
     TEAL, "registro.br"),
    ("🌐 Domínio .com", "Cloudflare Registrar",
     "US$ 10-12/ano\n≈ R$ 57-68/ano", "≈ R$ 5,00/mês",
     "Sem markup — só o custo real de\nregistro. DNS grátis incluso.",
     BLUE, "cloudflare.com/products/registrar"),
    ("🔒 SSL (HTTPS)", "Cloudflare Free",
     "GRATUITO", "R$ 0,00/mês",
     "Certificado HTTPS automático.\nRenovado automaticamente.",
     GREEN, "cloudflare.com/ssl"),
    ("📡 DNS Gerenciado", "Cloudflare Free",
     "GRATUITO", "R$ 0,00/mês",
     "DNS ultrarrápido com proteção\ncontra ataques DDoS inclusa.",
     GREEN, "cloudflare.com/dns"),
    ("☁ Nuvem (Always Free)", "Oracle OCI",
     "GRATUITO", "R$ 0,00/mês",
     "4 OCPU + 24 GB RAM + 200 GB disco.\nSuficiente para o sistema atual.",
     AMBER, "oracle.com/cloud/free"),
    ("☁ Nuvem (Pago básico)", "Oracle OCI",
     "≈ US$ 18-20/mês\n≈ R$ 102-114/mês", "≈ R$ 110/mês",
     "Se precisar de mais recursos.\nUpgrade simples sem contrato.",
     NAVY, "oracle.com/cloud/pricing"),
]

cw15 = Inches(3.9)
for i, (name, source, price, monthly, desc, color, url) in enumerate(cost_cards):
    row = i // 3
    col = i % 3
    l15 = Inches(0.4) + col * Inches(4.28)
    t15 = Inches(1.6) + row * Inches(2.8)
    rrect(s15, l15, t15, cw15, Inches(2.55), fill=WHITE, line=BORDER)
    rect(s15, l15, t15, cw15, Inches(0.07), fill=color)
    txt(s15, name, l15+Inches(0.15), t15+Inches(0.1), cw15-Inches(0.3), Inches(0.42),
        size=13, bold=True, color=color if color != MGRAY else DGRAY)
    txt(s15, f"Fonte: {source}", l15+Inches(0.15), t15+Inches(0.52), cw15-Inches(0.3),
        Inches(0.3), size=10, color=MGRAY, italic=True)
    bg_c = TEAL_BG if "GRATUITO" in price else RGBColor(0xFE, 0xF9, 0xE7)
    fc = GREEN if "GRATUITO" in price else AMBER
    rrect(s15, l15+Inches(0.1), t15+Inches(0.82), cw15-Inches(0.2), Inches(0.7),
          fill=bg_c, line=None)
    txt(s15, price, l15+Inches(0.15), t15+Inches(0.85), cw15-Inches(0.3), Inches(0.65),
        size=13, bold=True, color=fc, align=PP_ALIGN.CENTER)
    txt(s15, desc, l15+Inches(0.15), t15+Inches(1.6), cw15-Inches(0.3), Inches(0.7),
        size=10, color=DGRAY)
    txt(s15, url, l15+Inches(0.15), t15+Inches(2.3), cw15-Inches(0.3), Inches(0.25),
        size=9, color=MGRAY, italic=True)

section_label(s15, "Casa Assistencial · Pesquisa de Custos · Fonte: Cloudflare, Oracle, Registro.br — Maio/2026")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 16 — COMPARATIVO DE CUSTOS (TABELA)
# ══════════════════════════════════════════════════════════════════════════════
s16 = prs.slides.add_slide(BLANK)
add_bg(s16, LGRAY)
title_bar(s16, "Comparativo de Custos",
          "Visão consolidada de todos os serviços necessários · valores em reais")

rows16 = [
    ["Serviço", "Função", "Custo Mensal", "Custo Anual", "Observações"],
    ["Domínio .com.br\n(Registro.br)",   "Endereço do sistema",      "R$ 3,33",      "R$ 40,00",      "Renovação anual automática"],
    ["DNS\n(Cloudflare Free)",            "Roteamento do endereço",   "GRÁTIS",       "GRÁTIS",        "Plano gratuito para sempre"],
    ["Certificado HTTPS\n(Cloudflare)",   "Segurança e criptografia", "GRÁTIS",       "GRÁTIS",        "Renovação automática"],
    ["Servidor OCI\n(Always Free)",       "Hospedar o sistema",       "GRÁTIS",       "GRÁTIS",        "4 OCPU + 24GB RAM + 200GB"],
    ["TOTAL — Cenário Free",              "Sistema completo",         "≈ R$ 3,33",    "≈ R$ 40,00",    "✓ Recomendado para início"],
    ["Servidor OCI\n(Pago básico)",       "Mais capacidade",          "≈ R$ 110,00",  "≈ R$ 1.320,00", "Se o sistema crescer muito"],
    ["TOTAL — Cenário Pago",              "Sistema escalado",         "≈ R$ 113,33",  "≈ R$ 1.360,00", "Upgrade sem reescrever nada"],
]

col_widths16 = [
    Inches(2.4), Inches(2.6), Inches(1.9), Inches(1.9), Inches(4.1)
]
add_styled_table(s16, Inches(0.4), Inches(1.6), Inches(12.9), rows16, col_widths16)

# Nota
rrect(s16, Inches(0.4), Inches(6.7), Inches(12.9), Inches(0.5), fill=TEAL_BG, line=TEAL)
txt(s16, "* Cotação: US$ 1,00 = R$ 5,70 (mai/2026). Preços sujeitos a variação cambial. "
         "Fontes: Registro.br, Cloudflare, Oracle Cloud.",
    Inches(0.6), Inches(6.75), Inches(12.6), Inches(0.4), size=10, color=TEAL, italic=True)

section_label(s16, "Casa Assistencial · Comparativo de Custos")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 17 — CENÁRIO RECOMENDADO
# ══════════════════════════════════════════════════════════════════════════════
s17 = prs.slides.add_slide(BLANK)
add_bg(s17, LGRAY)
title_bar(s17, "Cenário Econômico Recomendado",
          "A opção de menor custo com máxima capacidade para o início do projeto")

# Destaque principal
rrect(s17, Inches(0.5), Inches(1.6), Inches(12.3), Inches(2.2), fill=TEAL, line=None)
txt(s17, "✅  Recomendação: OCI Always Free + Cloudflare Free + Registro.br",
    Inches(0.8), Inches(1.7), Inches(12.0), Inches(0.6),
    size=20, bold=True, color=WHITE)
txt(s17, "Custo estimado para o primeiro ano: R$ 40,00 (apenas o domínio .com.br)",
    Inches(0.8), Inches(2.3), Inches(12.0), Inches(0.5),
    size=16, color=WHITE)
txt(s17, "Servidor, DNS, HTTPS e banco de dados: completamente GRATUITOS",
    Inches(0.8), Inches(2.8), Inches(12.0), Inches(0.45),
    size=14, color=TEAL_BG, italic=True)

# Por que essa escolha?
txt(s17, "Por que essa é a melhor escolha para começar?",
    Inches(0.5), Inches(4.0), Inches(12), Inches(0.45),
    size=15, bold=True, color=NAVY)

reasons = [
    ("💰 Custo Mínimo",    "Começa com apenas R$ 40/ano. Sem surpresas na fatura.",                    TEAL),
    ("📈 Sem Compromisso", "Se o projeto crescer, basta fazer upgrade na OCI. Sem reescrever nada.",    BLUE),
    ("🔒 Segurança",       "HTTPS grátis e proteção DDoS do Cloudflare inclusos automaticamente.",      GREEN),
    ("⚡ Capacidade",      "4 OCPU e 24 GB RAM é mais que suficiente para dezenas de usuários.",        NAVY),
]
cw17 = Inches(3.0)
for i, (icon_t, desc_r, color) in enumerate(reasons):
    l17 = Inches(0.5) + i * Inches(3.2)
    rrect(s17, l17, Inches(4.55), cw17, Inches(2.1), fill=WHITE, line=BORDER)
    rect(s17, l17, Inches(4.55), cw17, Inches(0.07), fill=color)
    txt(s17, icon_t, l17+Inches(0.15), Inches(4.65), cw17-Inches(0.3), Inches(0.45),
        size=13, bold=True, color=color)
    txt(s17, desc_r, l17+Inches(0.15), Inches(5.1), cw17-Inches(0.3), Inches(1.3),
        size=11, color=DGRAY)

# Quando considerar upgrade?
rrect(s17, Inches(0.5), Inches(6.8), Inches(12.3), Inches(0.45), fill=RGBColor(0xFE, 0xF9, 0xE7), line=AMBER)
txt(s17, "📌  Considere upgrade quando: usuários simultâneos > 50 | banco de dados > 15 GB | necessidade de maior velocidade",
    Inches(0.7), Inches(6.85), Inches(12.0), Inches(0.38), size=11, color=RGBColor(0x7D,0x68,0x08))

section_label(s17, "Casa Assistencial · Cenário Recomendado")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 18 — ESCALABILIDADE
# ══════════════════════════════════════════════════════════════════════════════
s18 = prs.slides.add_slide(BLANK)
add_bg(s18, LGRAY)
title_bar(s18, "Escalabilidade — O Sistema Pode Crescer",
          "Sem reescrever nada: o sistema foi projetado para crescer junto com a instituição")

stages = [
    ("📌 Hoje\n(Início)",
     ["Sistema em uso", "1 unidade", "Até 20 usuários", "OCI Always Free", "Custo: R$ 40/ano"],
     TEAL, Inches(0.4)),
    ("📈 Fase 2\n(Crescimento)",
     ["Múltiplas unidades", "Até 100 usuários", "Mais dados e alertas", "OCI pago básico", "Custo: ≈ R$ 1.400/ano"],
     BLUE, Inches(3.6)),
    ("🚀 Fase 3\n(Expansão)",
     ["Integração com e-mail", "App mobile nativo", "Relatórios avançados", "Infraestrutura robusta", "Custo: sob demanda"],
     NAVY, Inches(6.8)),
    ("🌐 Fase 4\n(Escala)",
     ["Múltiplas instituições", "API pública", "Integrações externas", "Alta disponibilidade", "Modelo SaaS possível"],
     RGBColor(0x8E,0x44,0xAD), Inches(10.0)),
]

bw18 = Inches(3.0)
bh18 = Inches(4.5)
bt18 = Inches(2.1)
for (title18, items, color, l) in stages:
    rrect(s18, l, bt18, bw18, bh18, fill=color)
    txt(s18, title18, l, bt18+Inches(0.15), bw18, Inches(0.75),
        size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rect(s18, l+Inches(0.3), bt18+Inches(0.85), bw18-Inches(0.6), Inches(0.03), fill=WHITE)
    yo18 = bt18 + Inches(1.0)
    for item in items:
        txt(s18, f"• {item}", l+Inches(0.2), yo18, bw18-Inches(0.4), Inches(0.42),
            size=11, color=WHITE)
        yo18 += Inches(0.45)

# Setas de crescimento
for l_arr in [Inches(3.42), Inches(6.62), Inches(9.82)]:
    arrow_r(s18, l_arr, Inches(4.3), Inches(0.16), Inches(0.25), fill=AMBER)

# Nota importante
rrect(s18, Inches(0.4), Inches(6.8), Inches(12.5), Inches(0.45), fill=TEAL_BG, line=TEAL)
txt(s18, "✅  Cada fase é uma evolução natural. Nunca há necessidade de refazer o sistema — apenas expandir.",
    Inches(0.6), Inches(6.85), Inches(12.2), Inches(0.38), size=12, color=TEAL, bold=True)

section_label(s18, "Casa Assistencial · Escalabilidade")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 19 — BENEFÍCIOS PARA A INSTITUIÇÃO
# ══════════════════════════════════════════════════════════════════════════════
s19 = prs.slides.add_slide(BLANK)
add_bg(s19, LGRAY)
title_bar(s19, "Benefícios para a Instituição",
          "O impacto real do sistema no dia a dia da gestão")

benefits = [
    ("📋 Maior Controle", TEAL,
     ["Todas as informações de pacientes\ncentralizadas em um lugar",
      "Histórico completo de administrações\ne movimentações de estoque",
      "Visão em tempo real de tudo\nque acontece na instituição"]),
    ("🎯 Menos Erros", BLUE,
     ["Alertas automáticos evitam\nesquecimento de doses",
      "Registro obrigatório de cada\nadministração realizada",
      "Dupla verificação por\nprescrição e estoque"]),
    ("💰 Redução de Desperdícios", GREEN,
     ["Controle preciso de consumo\nevita pedidos em excesso",
      "Alertas de estoque mínimo\nevitam falta de insumos",
      "Rastreamento de perdas\ne descartes registrado"]),
    ("📊 Decisões Melhores", AMBER,
     ["Dashboard com indicadores\noperacionais em tempo real",
      "Histórico de consumo para\nplanejamento de compras",
      "Dados confiáveis para\nrelatórios e auditorias"]),
    ("🔒 Rastreabilidade Total", NAVY,
     ["Quem administrou, quando\ne em qual quantidade",
      "Log completo de alterações\ne movimentações",
      "Conformidade com normas\nassistenciais e sanitárias"]),
    ("📱 Acesso Moderno", RGBColor(0x8E,0x44,0xAD),
     ["Sistema acessível de qualquer\ndispositivo com internet",
      "Funciona offline (PWA)\npara uso sem conexão",
      "Interface intuitiva que\nnão exige treinamento longo"]),
]

cw19 = Inches(3.9)
for i, (title19, color, items) in enumerate(benefits):
    row = i // 3
    col = i % 3
    l19 = Inches(0.4) + col * Inches(4.28)
    t19 = Inches(1.6) + row * Inches(2.75)
    rrect(s19, l19, t19, cw19, Inches(2.5), fill=WHITE, line=BORDER)
    rect(s19, l19, t19, cw19, Inches(0.07), fill=color)
    txt(s19, title19, l19+Inches(0.15), t19+Inches(0.12), cw19-Inches(0.3), Inches(0.42),
        size=14, bold=True, color=color)
    yo19 = t19 + Inches(0.6)
    for item in items:
        txt(s19, f"• {item}", l19+Inches(0.18), yo19, cw19-Inches(0.36), Inches(0.6),
            size=10, color=DGRAY)
        yo19 += Inches(0.6)

section_label(s19, "Casa Assistencial · Benefícios")

# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 20 — CONCLUSÃO
# ══════════════════════════════════════════════════════════════════════════════
s20 = prs.slides.add_slide(BLANK)
add_bg(s20, NAVY)
rect(s20, 0, 0, Inches(0.75), H, fill=TEAL)

# Círculos decorativos
oval(s20, Inches(10.5), Inches(-1.5), Inches(5), Inches(5), fill=BLUE)
oval(s20, Inches(11.2), Inches(-0.8), Inches(3.2), Inches(3.2), fill=TEAL)
oval(s20, Inches(9.5), Inches(5.5), Inches(2.5), Inches(2.5), fill=RGBColor(0x1A,0x2E,0x6A))

txt(s20, "Conclusão", Inches(1.2), Inches(0.7), Inches(9), Inches(0.65),
    size=13, color=TEAL_LT)
txt(s20, "Um sistema moderno,\nacessível e viável.",
    Inches(1.2), Inches(1.3), Inches(10), Inches(1.8),
    size=42, bold=True, color=WHITE)
rect(s20, Inches(1.2), Inches(3.15), Inches(7), Inches(0.045), fill=TEAL_LT)

# 3 resumos
summary = [
    ("🏗 Arquitetura Sólida",
     "Frontend React + Backend FastAPI + PostgreSQL\nTestado, documentado e pronto para produção"),
    ("☁ Infraestrutura Inteligente",
     "OCI Always Free + Cloudflare + Registro.br\nCusto inicial de apenas R$ 40/ano"),
    ("📈 Pronto para Crescer",
     "Escalabilidade sem reescrever\nExpande conforme a demanda aumenta"),
]
yo20 = Inches(3.4)
for (title20, desc20) in summary:
    rrect(s20, Inches(1.2), yo20, Inches(9.8), Inches(0.95), fill=RGBColor(0x22,0x3A,0x5E), line=None)
    rect(s20, Inches(1.2), yo20, Inches(0.07), Inches(0.95), fill=TEAL)
    txt(s20, title20, Inches(1.45), yo20+Inches(0.07), Inches(9.5), Inches(0.38),
        size=14, bold=True, color=WHITE)
    txt(s20, desc20, Inches(1.45), yo20+Inches(0.47), Inches(9.5), Inches(0.4),
        size=11, color=MGRAY)
    yo20 += Inches(1.05)

# CTA
rrect(s20, Inches(1.2), Inches(6.6), Inches(9.8), Inches(0.65), fill=TEAL, line=None)
txt(s20, "🚀  O sistema está em desenvolvimento e pronto para entrar em produção.",
    Inches(1.5), Inches(6.68), Inches(9.5), Inches(0.5),
    size=14, bold=True, color=WHITE)

txt(s20, "Casa Assistencial  ·  2026  ·  Plataforma de Gestão Assistencial em Nuvem",
    Inches(1.2), Inches(7.2), Inches(10), Inches(0.3),
    size=10, color=MGRAY)

# ── Salvar ─────────────────────────────────────────────────────────────────────
output_path = r"C:\Users\User\Desktop\Chris\Ed\MainAPI-main\Apresentacao_CasaAssistencial.pptx"
prs.save(output_path)
print(f"Apresentação gerada: {output_path}")
print(f"Total de slides: {len(prs.slides)}")
